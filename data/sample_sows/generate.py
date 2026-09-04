"""Render the synthetic SOWs in content.py to TXT, DOCX, and PDF.

Run from the repo root:
    .venv/Scripts/python data/sample_sows/generate.py
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from content import ALL_SOWS

OUT_DIR = Path(__file__).parent


def render_txt(sow: dict, path: Path) -> None:
    lines = [sow["doc_title"], "=" * len(sow["doc_title"]), ""]
    for section in sow["sections"]:
        lines += [section["title"], "-" * len(section["title"]), ""]
        for p in section["paragraphs"]:
            lines += [p, ""]
        if table := section.get("table"):
            widths = [max(len(row[i]) for row in table) for i in range(len(table[0]))]
            for r_i, row in enumerate(table):
                lines.append("  ".join(cell.ljust(w) for cell, w in zip(row, widths)))
                if r_i == 0:
                    lines.append("  ".join("-" * w for w in widths))
            lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def render_docx(sow: dict, path: Path) -> None:
    doc = Document()
    doc.add_heading(sow["doc_title"], level=0)
    for section in sow["sections"]:
        doc.add_heading(section["title"], level=1)
        for p in section["paragraphs"]:
            doc.add_paragraph(p)
        if table := section.get("table"):
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            t.style = "Light Grid Accent 1"
            for r_i, row in enumerate(table):
                for c_i, cell in enumerate(row):
                    t.rows[r_i].cells[c_i].text = cell
    doc.save(str(path))


def render_pdf(sow: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10.5, leading=14)
    story = [Paragraph(sow["doc_title"], styles["Title"]), Spacer(1, 0.5 * cm)]
    for section in sow["sections"]:
        story.append(Paragraph(section["title"], styles["Heading1"]))
        for p in section["paragraphs"]:
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 0.2 * cm))
        if table := section.get("table"):
            t = Table([[Paragraph(c, body) for c in row] for row in table], hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 0.3 * cm))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def main() -> None:
    for slug, sow in ALL_SOWS.items():
        render_txt(sow, OUT_DIR / f"{slug}.txt")
        render_docx(sow, OUT_DIR / f"{slug}.docx")
        render_pdf(sow, OUT_DIR / f"{slug}.pdf")
        print(f"rendered {slug} -> txt, docx, pdf")


if __name__ == "__main__":
    main()
